from pathlib import Path
import collections
import copy
import json
import openpyxl

DIR = Path(__file__).resolve().parent
original = openpyxl.load_workbook(DIR/'grammar_tasks.original.xlsx')
updated = openpyxl.load_workbook(DIR.parent.parent/'collections/grammar_tasks.xlsx')
x,y = original.worksheets[0],updated.worksheets[0]
patches = {p['cell']:p for p in json.loads((DIR/'patches.json').read_text(encoding='utf8'))}
value_errors=[]
style_diffs=collections.Counter()
examples={}
changed=[]
assert original.sheetnames==updated.sheetnames
assert (x.max_row,x.max_column)==(y.max_row,y.max_column)
for row in x:
    for before in row:
        after=y[before.coordinate]
        patch=patches.get(before.coordinate)
        expected=patch['value'] if patch else before.value
        if patch and before.value!=patch['old']:
            value_errors.append((before.coordinate,'unexpected original',before.value,patch['old']))
        if after.value!=expected:
            value_errors.append((before.coordinate,'unexpected final',expected,after.value))
        if after.value!=before.value: changed.append(before.coordinate)
        for name in ('font','fill','border','alignment','number_format','protection'):
            a,b=copy.copy(getattr(before,name)),copy.copy(getattr(after,name))
            if a!=b:
                style_diffs[name]+=1
                examples.setdefault(name,[before.coordinate,str(a),str(b)])
properties=dict(
    original_freeze=x.freeze_panes,final_freeze=y.freeze_panes,
    original_filter=x.auto_filter.ref,final_filter=y.auto_filter.ref,
    original_tables=list(x.tables),final_tables=list(y.tables),
    original_merges=[str(a) for a in x.merged_cells],final_merges=[str(a) for a in y.merged_cells],
    original_widths={k:v.width for k,v in x.column_dimensions.items()},
    final_widths={k:v.width for k,v in y.column_dimensions.items()},
    original_gridlines=x.sheet_view.showGridLines,final_gridlines=y.sheet_view.showGridLines,
)
result=dict(changed_cells=len(changed),changed_tasks=len({patches[c]['id'] for c in changed}),
    value_errors=value_errors,style_diffs=dict(style_diffs),style_examples=examples,properties=properties)
(DIR/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False,indent=2))
assert not value_errors
